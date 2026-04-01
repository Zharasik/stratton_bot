import os
import logging
from docx import Document
from datetime import datetime

logger = logging.getLogger(__name__)
TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "data", "nda_template.docx")

MONTHS = {1:"января",2:"февраля",3:"марта",4:"апреля",5:"мая",6:"июня",7:"июля",8:"августа",9:"сентября",10:"октября",11:"ноября",12:"декабря"}


def generate_nda(data: dict, output_path: str) -> bool:
    try:
        doc = Document(TEMPLATE)
        now = datetime.now()
        name = f"{data.get('last_name','')} {data.get('first_name','')} {data.get('middle_name','')}".strip()
        reps = {
            "ФИО": name,
            "Адрес по прописке и фактический:": f"Адрес по прописке и фактический:\n{data.get('address','')}",
            "Мобильный:": f"Мобильный: {data.get('phone','')}",
            "ИИН:": f"ИИН: {data.get('iin','')}",
            "Email:": f"Email: {data.get('email','')}",
            "«__» _______ 20__ года": f"«{now.day}» {MONTHS.get(now.month,'')} {now.year} года",
        }
        for p in doc.paragraphs:
            for old, new in reps.items():
                if old in p.text:
                    for run in p.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for old, new in reps.items():
                            if old in p.text:
                                for run in p.runs:
                                    if old in run.text:
                                        run.text = run.text.replace(old, new)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return True
    except Exception as e:
        logger.error("NDA gen error: %s", e)
        return False
