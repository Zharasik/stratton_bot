import os
import logging
from datetime import datetime, timedelta
from typing import Optional
from zoneinfo import ZoneInfo

from dto import UserDTO, SlotDTO, TestingDTO, NdaDTO
from repositories import (
    UserRepository, SlotRepository, TestingRepository,
    SubmissionRepository, NdaRepository, StatsRepository,
)
from settings import AppSettings

logger = logging.getLogger(__name__)

TASK_VARIANTS = {
    1: "📝 Вариант 1 — To-Do бот",
    2: "💰 Вариант 2 — Бот учета расходов",
    3: "🎮 Вариант 3 — Угадай число",
    4: "📊 Вариант 4 — Бот-опросник",
    5: "📅 Вариант 5 — Бот напоминаний",
    6: "🎲 Вариант 6 — Казино-бот",
    7: "📚 Вариант 7 — Бот для изучения слов",
    8: "🧾 Вариант 8 — Генератор QR-кодов",
    9: "🏆 Вариант 9 — Викторина",
    10: "🤝 Вариант 10 — Бот поиска собеседника",
}

TASK_HEADER = (
    "⏱️ Время: 4 часа\n"
    "✅ Задание на проверке после отправки видео/ссылки.\n"
    "🐍 aiogram 3\n\n"
)
TASK_FOOTER = "\n📹 Видео до 60 сек, до 10 МБ"

TASK_BODIES = {
    1: "📝 <b>Задание:</b>\nTo-Do бот:\n- Добавление/удаление задач\n- Просмотр списка\n- Отметка выполненных\n\nТех: aiogram 3.x, БД, FSM, обработка ошибок, логирование\nДоп: Фильтрация, документация, README.md, requirements.txt",
    2: "📝 <b>Задание:</b>\nУчёт расходов:\n- Добавление (сумма + категория)\n- Статистика\n- Все расходы\n\nТех: aiogram 3.x, БД, валидация, обработка ошибок, логирование\nДоп: Группировка, документация, README.md, requirements.txt",
    3: "📝 <b>Задание:</b>\nУгадай число:\n- Загадывание числа\n- Подсказки больше/меньше\n- Подсчет попыток\n\nТех: aiogram 3.x, БД, FSM, обработка ошибок, логирование\nДоп: Рекорды, документация, README.md, requirements.txt",
    4: "📝 <b>Задание:</b>\nОпросник:\n- Создание опроса\n- Голосование\n- Результаты\n\nТех: aiogram 3.x, БД, Inline кнопки, обработка ошибок, логирование\nДоп: Один голос, документация, README.md, requirements.txt",
    5: "📝 <b>Задание:</b>\nНапоминания:\n- Добавление\n- Время\n- Уведомление\n\nТех: aiogram 3.x, БД, планировщик, обработка ошибок, логирование\nДоп: Повторяющиеся, документация, README.md, requirements.txt",
    6: "📝 <b>Задание:</b>\nКазино:\n- Баланс\n- Ставки\n- Орёл/решка\n\nТех: aiogram 3.x, БД, валидация ставок, обработка ошибок, логирование\nДоп: Лидерборд, документация, README.md, requirements.txt",
    7: "📝 <b>Задание:</b>\nИзучение слов:\n- Добавление\n- Тестирование\n- Подсчет\n\nТех: aiogram 3.x, БД, FSM, обработка ошибок, логирование\nДоп: Случайный порядок, документация, README.md, requirements.txt",
    8: "📝 <b>Задание:</b>\nQR-коды:\n- Ввод текста/ссылки\n- Генерация QR\n- Отправка картинки\n\nТех: aiogram 3.x, файлы, обработка ошибок, логирование\nДоп: История, документация, README.md, requirements.txt",
    9: "📝 <b>Задание:</b>\nВикторина:\n- Вопросы с вариантами\n- Подсчет очков\n- Результат\n\nТех: aiogram 3.x, БД, Inline, обработка ошибок, логирование\nДоп: Таймер, документация, README.md, requirements.txt",
    10: "📝 <b>Задание:</b>\nПоиск собеседника:\n- Регистрация\n- Поиск\n- Переписка через бота\n\nТех: aiogram 3.x, БД, состояния, обработка ошибок, логирование\nДоп: Очередь, документация, README.md, requirements.txt",
}


class TaskService:
    @staticmethod
    def get_variant_number(user_id: int) -> int:
        return (user_id % len(TASK_VARIANTS)) + 1

    @staticmethod
    def get_task_text(user_id: int) -> str:
        v = TaskService.get_variant_number(user_id)
        return f"{TASK_VARIANTS[v]}\n\n{TASK_HEADER}{TASK_BODIES[v]}{TASK_FOOTER}"

    @staticmethod
    def is_within_window(slot_date: str, slot_start: str, task_hours: int) -> bool:
        now = datetime.now()
        start = datetime.strptime(f"{slot_date} {slot_start}", "%Y-%m-%d %H:%M")
        return start <= now <= start + timedelta(hours=task_hours)

    @staticmethod
    def get_deadline(slot_date: str, slot_start: str, task_hours: int) -> str:
        start = datetime.strptime(f"{slot_date} {slot_start}", "%Y-%m-%d %H:%M")
        return (start + timedelta(hours=task_hours)).strftime("%H:%M")

    @staticmethod
    def get_remaining_minutes(slot_date: str, slot_start: str, task_hours: int) -> int:
        start = datetime.strptime(f"{slot_date} {slot_start}", "%Y-%m-%d %H:%M")
        end = start + timedelta(hours=task_hours)
        remaining = end - datetime.now()
        return max(0, int(remaining.total_seconds() // 60))


class NdaDocumentService:
    NDA_TEMPLATE = os.path.join(os.path.dirname(__file__), "static", "nda_template.docx")
    MONTHS = {1: "января", 2: "февраля", 3: "марта", 4: "апреля", 5: "мая", 6: "июня",
              7: "июля", 8: "августа", 9: "сентября", 10: "октября", 11: "ноября", 12: "декабря"}

    @classmethod
    def generate(cls, data: dict, output_path: str) -> bool:
        try:
            from docx import Document
            doc = Document(cls.NDA_TEMPLATE)
            now = datetime.now()
            name = f"{data.get('last_name', '')} {data.get('first_name', '')} {data.get('middle_name', '')}".strip()
            replacements = {
                "ФИО": name,
                "Адрес по прописке и фактический:": f"Адрес по прописке и фактический:\n{data.get('address', '')}",
                "Мобильный:": f"Мобильный: {data.get('phone', '')}",
                "ИИН:": f"ИИН: {data.get('iin', '')}",
                "Email:": f"Email: {data.get('email', '')}",
                "«__» _______ 20__ года": f"«{now.day}» {cls.MONTHS.get(now.month, '')} {now.year} года",
            }
            for paragraph in doc.paragraphs:
                cls._replace_in_paragraph(paragraph, replacements)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            cls._replace_in_paragraph(paragraph, replacements)
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            doc.save(output_path)
            return True
        except Exception as e:
            logger.error("NDA generation failed: %s", e)
            return False

    @staticmethod
    def _replace_in_paragraph(paragraph, replacements: dict):
        for old, new in replacements.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
