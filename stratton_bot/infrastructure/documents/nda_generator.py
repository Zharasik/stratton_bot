from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

from stratton_bot.domain.entities import NdaRecordData


class NDADocumentGenerator:
    MONTHS = {
        1: "января",
        2: "февраля",
        3: "марта",
        4: "апреля",
        5: "мая",
        6: "июня",
        7: "июля",
        8: "августа",
        9: "сентября",
        10: "октября",
        11: "ноября",
        12: "декабря",
    }

    def __init__(self, template_path: Path) -> None:
        self._template_path = template_path

    def generate(self, nda: NdaRecordData, output_path: Path) -> Path:
        document = Document(self._template_path)
        now = datetime.now()
        full_name = f"{nda.last_name} {nda.first_name} {nda.middle_name}".strip()
        replacements = {
            "ФИО": full_name,
            "Адрес по прописке и фактический:": f"Адрес по прописке и фактический:\n{nda.address}",
            "Мобильный:": f"Мобильный: {nda.phone}",
            "ИИН:": f"ИИН: {nda.iin}",
            "Email:": f"Email: {nda.email}",
            "«__» _______ 20__ года": f"«{now.day}» {self.MONTHS.get(now.month, '')} {now.year} года",
        }
        for paragraph in document.paragraphs:
            self._replace_in_paragraph(paragraph, replacements)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph, replacements)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return output_path

    @staticmethod
    def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
        for old, new in replacements.items():
            if old not in paragraph.text:
                continue
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
